"""
Reports Service
Handles report generation logic and data aggregation
"""
from motor.motor_asyncio import AsyncIOMotorDatabase
from typing import Optional, Dict, List, Tuple
import pandas as pd
import io
from datetime import datetime
import logging

logger = logging.getLogger(__name__)

# Check if openpyxl is available for Excel export
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.utils.dataframe import dataframe_to_rows
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    logger.warning("openpyxl not available. Excel export with formatting will be limited.")


class ReportsService:
    def __init__(self, db: AsyncIOMotorDatabase):
        self.db = db
        self.collection = db.business_data

    def _parse_filter_string(self, filter_str: Optional[str]) -> Optional[List[str]]:
        """Parse comma-separated filter string into list"""
        if not filter_str:
            return None
        return [item.strip() for item in filter_str.split(',') if item.strip()]

    async def _get_filtered_data(
        self,
        years: Optional[str] = None,
        months: Optional[str] = None,
        businesses: Optional[str] = None,
        channels: Optional[str] = None,
        brands: Optional[str] = None,
        categories: Optional[str] = None,
        customers: Optional[str] = None
    ) -> pd.DataFrame:
        """Get filtered data from MongoDB"""
        query = {}
        
        # Build query filters
        if years:
            year_list = self._parse_filter_string(years)
            if year_list:
                query['Year'] = {'$in': [int(y) for y in year_list]}
        
        if months:
            month_list = self._parse_filter_string(months)
            if month_list:
                query['Month'] = {'$in': month_list}
        
        if businesses:
            business_list = self._parse_filter_string(businesses)
            if business_list:
                query['Business'] = {'$in': business_list}
        
        if channels:
            channel_list = self._parse_filter_string(channels)
            if channel_list:
                query['Channel'] = {'$in': channel_list}
        
        if brands:
            brand_list = self._parse_filter_string(brands)
            if brand_list:
                query['Brand'] = {'$in': brand_list}
        
        if categories:
            category_list = self._parse_filter_string(categories)
            if category_list:
                query['Category'] = {'$in': category_list}
        
        if customers:
            customer_list = self._parse_filter_string(customers)
            if customer_list:
                query['Customer'] = {'$in': customer_list}
        
        # Fetch data from MongoDB
        cursor = self.collection.find(query)
        data = await cursor.to_list(length=None)
        
        if not data:
            return pd.DataFrame()
        
        # Convert to DataFrame
        df = pd.DataFrame(data)
        
        # Remove MongoDB _id field
        if '_id' in df.columns:
            df = df.drop('_id', axis=1)
        
        return df

    def _format_excel_worksheet(self, ws, df: pd.DataFrame, title: str):
        """Apply formatting to Excel worksheet"""
        if not EXCEL_AVAILABLE:
            return
        
        # Title formatting
        ws['A1'] = title
        ws['A1'].font = Font(size=16, bold=True, color="FFFFFF")
        ws['A1'].fill = PatternFill(start_color="184464", end_color="184464", fill_type="solid")
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells(f'A1:{openpyxl.utils.get_column_letter(len(df.columns))}1')
        ws.row_dimensions[1].height = 30
        
        # Header formatting
        header_fill = PatternFill(start_color="D9E9F7", end_color="D9E9F7", fill_type="solid")
        header_font = Font(bold=True, size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for cell in ws[3]:  # Headers are in row 3
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Data formatting
        for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal='left', vertical='center')
        
        # Auto-adjust column widths
        for column_cells in ws.columns:
            max_length = 0
            column_letter = None
            
            for cell in column_cells:
                try:
                    # Skip merged cells
                    if hasattr(cell, 'column_letter'):
                        if column_letter is None:
                            column_letter = cell.column_letter
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            
            if column_letter:
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

    def _create_excel_with_formatting(self, data_dict: Dict[str, pd.DataFrame], report_title: str) -> bytes:
        """Create Excel file with multiple sheets and formatting"""
        output = io.BytesIO()
        
        if EXCEL_AVAILABLE:
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                for sheet_name, df in data_dict.items():
                    if df.empty:
                        continue
                    
                    # Write DataFrame starting from row 3 (row 1 for title, row 2 empty, row 3 headers)
                    df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=2)
                    
                    # Get worksheet and apply formatting
                    ws = writer.sheets[sheet_name]
                    self._format_excel_worksheet(ws, df, f"{report_title} - {sheet_name}")
        else:
            # Fallback to basic Excel export
            with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
                for sheet_name, df in data_dict.items():
                    if df.empty:
                        continue
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        output.seek(0)
        return output.getvalue()

    async def generate_custom_report(
        self,
        years: Optional[str] = None,
        months: Optional[str] = None,
        businesses: Optional[str] = None,
        channels: Optional[str] = None,
        brands: Optional[str] = None,
        categories: Optional[str] = None,
        format: str = "excel"
    ) -> Tuple[bytes, str, str]:
        """Generate custom report based on filters"""
        try:
            # Get filtered data
            df = await self._get_filtered_data(
                years=years,
                months=months,
                businesses=businesses,
                channels=channels,
                brands=brands,
                categories=categories
            )
            
            if df.empty:
                raise ValueError("No data found for the selected filters")
            
            # Generate summary statistics
            summary_df = self._generate_summary_statistics(df)
            
            # Prepare data dictionary
            data_dict = {
                'Summary': summary_df,
                'Detailed Data': df
            }
            
            # Generate filename
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            
            if format.lower() == "csv":
                output = io.BytesIO()
                df.to_csv(output, index=False)
                output.seek(0)
                return output.getvalue(), f"custom_report_{timestamp}.csv", "text/csv"
            else:
                file_data = self._create_excel_with_formatting(data_dict, "Custom Report")
                return file_data, f"custom_report_{timestamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            
        except Exception as e:
            logger.error(f"Error generating custom report: {str(e)}")
            raise

    def _ensure_numeric(self, df: pd.DataFrame, columns: list) -> pd.DataFrame:
        """Ensure specified columns are numeric"""
        for col in columns:
            if col in df.columns:
                df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0)
        return df

    def _generate_summary_statistics(self, df: pd.DataFrame) -> pd.DataFrame:
        """Generate summary statistics from data"""
        if df.empty:
            return pd.DataFrame({'Metric': [], 'Value': []})
        
        # Ensure numeric columns (using actual MongoDB column names)
        numeric_cols = ['Revenue', 'Gross_Profit', 'Units']
        df = self._ensure_numeric(df, numeric_cols)
        
        summary_data = []
        
        # Revenue metrics
        if 'Revenue' in df.columns:
            total_sales = df['Revenue'].sum()
            avg_sales = df['Revenue'].mean()
            summary_data.append({'Metric': 'Total Revenue', 'Value': f"${total_sales:,.2f}"})
            summary_data.append({'Metric': 'Average Revenue per Transaction', 'Value': f"${avg_sales:,.2f}"})
        
        # Profit metrics
        if 'Gross_Profit' in df.columns:
            total_profit = df['Gross_Profit'].sum()
            summary_data.append({'Metric': 'Total Profit', 'Value': f"${total_profit:,.2f}"})
            
            if 'Revenue' in df.columns and df['Revenue'].sum() > 0:
                margin = (total_profit / df['Revenue'].sum() * 100)
                summary_data.append({'Metric': 'Average Profit Margin', 'Value': f"{margin:.2f}%"})
        
        # Units metrics
        if 'Units' in df.columns:
            total_units = df['Units'].sum()
            summary_data.append({'Metric': 'Total Units Sold', 'Value': f"{total_units:,.0f}"})
        
        # Transaction count
        summary_data.append({'Metric': 'Total Transactions', 'Value': f"{len(df):,}"})
        
        return pd.DataFrame(summary_data)

    async def generate_executive_summary_report(
        self,
        years: Optional[str] = None,
        months: Optional[str] = None,
        businesses: Optional[str] = None
    ) -> Tuple[bytes, str, str]:
        """Generate Executive Summary Report"""
        try:
            df = await self._get_filtered_data(years=years, months=months, businesses=businesses)
            
            if df.empty:
                raise ValueError("No data found for executive summary. Please adjust your filters or ensure data exists in the database.")
            
            # Ensure numeric columns
            numeric_cols = ['Revenue', 'Gross_Profit', 'Units']
            df = self._ensure_numeric(df, numeric_cols)
            
            # Overall KPIs
            total_sales = df['Revenue'].sum() if 'Revenue' in df.columns else 0
            total_profit = df['Gross_Profit'].sum() if 'Gross_Profit' in df.columns else 0
            profit_margin = (total_profit / total_sales * 100) if total_sales > 0 else 0
            total_units = df['Units'].sum() if 'Units' in df.columns else 0
            avg_transaction = df['Revenue'].mean() if 'Revenue' in df.columns else 0
            
            kpi_data = {
                'Metric': ['Total Revenue', 'Total Profit', 'Profit Margin %', 'Total Units', 'Transactions', 'Avg Transaction Value'],
                'Value': [
                    f"${total_sales:,.2f}",
                    f"${total_profit:,.2f}",
                    f"{profit_margin:.2f}%",
                    f"{total_units:,.0f}",
                    f"{len(df):,}",
                    f"${avg_transaction:,.2f}"
                ]
            }
            kpi_df = pd.DataFrame(kpi_data)
            
            # Top performers
            if 'Customer' in df.columns and 'Revenue' in df.columns and not df.empty:
                top_customers = df.groupby('Customer')['Revenue'].sum().sort_values(ascending=False).head(10).reset_index()
                top_customers.columns = ['Customer', 'Revenue']
                top_customers['Revenue'] = top_customers['Revenue'].apply(lambda x: f"${x:,.2f}")
            else:
                top_customers = pd.DataFrame({'Customer': [], 'Revenue': []})
            
            if 'Brand' in df.columns and 'Revenue' in df.columns and not df.empty:
                top_brands = df.groupby('Brand')['Revenue'].sum().sort_values(ascending=False).head(10).reset_index()
                top_brands.columns = ['Brand', 'Revenue']
                top_brands['Revenue'] = top_brands['Revenue'].apply(lambda x: f"${x:,.2f}")
            else:
                top_brands = pd.DataFrame({'Brand': [], 'Revenue': []})
            
            # Year over year
            if 'Year' in df.columns and not df.empty:
                yoy_df = df.groupby('Year').agg({
                    'Revenue': 'sum',
                    'Gross_Profit': 'sum',
                    'Units': 'sum'
                }).reset_index()
                yoy_df.columns = ['Year', 'Revenue', 'Gross Profit', 'Units']
                yoy_df['Revenue'] = yoy_df['Revenue'].apply(lambda x: f"${x:,.2f}")
                yoy_df['Gross Profit'] = yoy_df['Gross Profit'].apply(lambda x: f"${x:,.2f}")
                yoy_df['Units'] = yoy_df['Units'].apply(lambda x: f"{x:,.0f}")
            else:
                yoy_df = pd.DataFrame({'Year': [], 'Revenue': [], 'Gross Profit': [], 'Units': []})
            
            data_dict = {
                'Executive KPIs': kpi_df,
                'Top 10 Customers': top_customers,
                'Top 10 Brands': top_brands,
                'Year over Year': yoy_df
            }
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_data = self._create_excel_with_formatting(data_dict, "Executive Summary")
            return file_data, f"executive_summary_{timestamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            
        except Exception as e:
            logger.error(f"Error generating executive summary: {str(e)}")
            raise

    async def generate_customer_performance_report(
        self,
        years: Optional[str] = None,
        months: Optional[str] = None,
        businesses: Optional[str] = None,
        channels: Optional[str] = None
    ) -> Tuple[bytes, str, str]:
        """Generate Customer Performance Report"""
        try:
            df = await self._get_filtered_data(years=years, months=months, businesses=businesses, channels=channels)
            
            if df.empty:
                raise ValueError("No data found for customer performance")
            
            # Customer analysis
            if 'Customer' in df.columns:
                customer_summary = df.groupby('Customer').agg({
                    'Revenue': 'sum',
                    'Gross_Profit': 'sum',
                    'Units': 'sum'
                }).reset_index()
                customer_summary.columns = ['Customer', 'Total Revenue', 'Total Profit', 'Units Sold']
                
                # Count transactions per customer
                transaction_counts = df.groupby('Customer').size().reset_index(name='Transactions')
                customer_summary = customer_summary.merge(transaction_counts, on='Customer')
                
                customer_summary['Avg Transaction Value'] = customer_summary['Total Revenue'] / customer_summary['Transactions']
                customer_summary['Profit Margin %'] = (customer_summary['Total Profit'] / customer_summary['Total Revenue'] * 100).round(2)
                customer_summary = customer_summary.sort_values('Total Revenue', ascending=False)
            else:
                customer_summary = pd.DataFrame()
            
            # Channel performance
            channel_df = df.groupby('Channel').agg({
                'Revenue': 'sum',
                'Gross_Profit': 'sum',
                'Units': 'sum'
            }).reset_index() if 'Channel' in df.columns else pd.DataFrame()
            
            if not channel_df.empty:
                channel_df.columns = ['Channel', 'Revenue', 'Profit', 'Units']
            
            # Customer by business
            customer_business = df.groupby(['Customer', 'Business'])['Revenue'].sum().reset_index() if 'Customer' in df.columns and 'Business' in df.columns else pd.DataFrame()
            customer_business = customer_business.sort_values('Revenue', ascending=False) if not customer_business.empty else customer_business
            if not customer_business.empty:
                customer_business.columns = ['Customer', 'Business', 'Revenue']
            
            data_dict = {
                'Customer Summary': customer_summary,
                'Channel Performance': channel_df,
                'Customer by Business': customer_business
            }
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_data = self._create_excel_with_formatting(data_dict, "Customer Performance")
            return file_data, f"customer_performance_{timestamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            
        except Exception as e:
            logger.error(f"Error generating customer performance report: {str(e)}")
            raise

    async def generate_brand_analysis_report(
        self,
        years: Optional[str] = None,
        months: Optional[str] = None,
        businesses: Optional[str] = None,
        brands: Optional[str] = None
    ) -> Tuple[bytes, str, str]:
        """Generate Brand Analysis Report"""
        try:
            df = await self._get_filtered_data(years=years, months=months, businesses=businesses, brands=brands)
            
            if df.empty:
                raise ValueError("No data found for brand analysis")
            
            # Brand summary
            if 'Brand' in df.columns:
                brand_summary = df.groupby('Brand').agg({
                    'Revenue': 'sum',
                    'Gross_Profit': 'sum',
                    'Units': 'sum'
                }).reset_index()
                brand_summary.columns = ['Brand', 'Revenue', 'Profit', 'Units']
                
                # Count transactions per brand
                transaction_counts = df.groupby('Brand').size().reset_index(name='Transactions')
                brand_summary = brand_summary.merge(transaction_counts, on='Brand')
                
                brand_summary['Profit Margin %'] = (brand_summary['Profit'] / brand_summary['Revenue'] * 100).round(2)
                brand_summary = brand_summary.sort_values('Revenue', ascending=False)
            else:
                brand_summary = pd.DataFrame()
            
            # Brand by category
            brand_category = df.groupby(['Brand', 'Category']).agg({
                'Revenue': 'sum',
                'Units': 'sum'
            }).reset_index() if 'Brand' in df.columns and 'Category' in df.columns else pd.DataFrame()
            
            if not brand_category.empty:
                brand_category.columns = ['Brand', 'Category', 'Revenue', 'Units']
                brand_category = brand_category.sort_values('Revenue', ascending=False)
            
            # Brand by channel
            brand_channel = df.groupby(['Brand', 'Channel'])['Revenue'].sum().reset_index() if 'Brand' in df.columns and 'Channel' in df.columns else pd.DataFrame()
            if not brand_channel.empty:
                brand_channel.columns = ['Brand', 'Channel', 'Revenue']
                brand_channel = brand_channel.sort_values('Revenue', ascending=False)
            
            data_dict = {
                'Brand Summary': brand_summary,
                'Brand by Category': brand_category,
                'Brand by Channel': brand_channel
            }
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_data = self._create_excel_with_formatting(data_dict, "Brand Analysis")
            return file_data, f"brand_analysis_{timestamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            
        except Exception as e:
            logger.error(f"Error generating brand analysis report: {str(e)}")
            raise

    async def generate_category_insights_report(
        self,
        years: Optional[str] = None,
        months: Optional[str] = None,
        businesses: Optional[str] = None,
        categories: Optional[str] = None
    ) -> Tuple[bytes, str, str]:
        """Generate Category Insights Report"""
        try:
            df = await self._get_filtered_data(years=years, months=months, businesses=businesses, categories=categories)
            
            if df.empty:
                raise ValueError("No data found for category insights")
            
            # Category summary
            category_summary = df.groupby('Category').agg({
                'Revenue': 'sum',
                'Gross_Profit': 'sum',
                'Units': 'sum'
            }).reset_index() if 'Category' in df.columns else pd.DataFrame()
            
            if not category_summary.empty:
                category_summary.columns = ['Category', 'Revenue', 'Profit', 'Units']
                category_summary['Profit Margin %'] = (category_summary['Profit'] / category_summary['Revenue'] * 100).round(2)
                category_summary = category_summary.sort_values('Revenue', ascending=False)
            
            # Category by business
            category_business = df.groupby(['Category', 'Business'])['Revenue'].sum().reset_index() if 'Category' in df.columns and 'Business' in df.columns else pd.DataFrame()
            if not category_business.empty:
                category_business.columns = ['Category', 'Business', 'Revenue']
                category_business = category_business.sort_values('Revenue', ascending=False)
            
            # Category trends by year
            category_year = df.groupby(['Category', 'Year'])['Revenue'].sum().reset_index() if 'Category' in df.columns and 'Year' in df.columns else pd.DataFrame()
            if not category_year.empty:
                category_year.columns = ['Category', 'Year', 'Revenue']
            
            data_dict = {
                'Category Summary': category_summary,
                'Category by Business': category_business,
                'Category Trends': category_year
            }
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_data = self._create_excel_with_formatting(data_dict, "Category Insights")
            return file_data, f"category_insights_{timestamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            
        except Exception as e:
            logger.error(f"Error generating category insights report: {str(e)}")
            raise

    async def generate_yoy_comparison_report(
        self,
        years: Optional[str] = None,
        businesses: Optional[str] = None
    ) -> Tuple[bytes, str, str]:
        """Generate Year-over-Year Comparison Report"""
        try:
            df = await self._get_filtered_data(years=years, businesses=businesses)
            
            if df.empty:
                raise ValueError("No data found for YoY comparison")
            
            # Year over year summary
            if 'Year' in df.columns:
                yoy_summary = df.groupby('Year').agg({
                    'Revenue': 'sum',
                    'Gross_Profit': 'sum',
                    'Units': 'sum'
                }).reset_index()
                yoy_summary.columns = ['Year', 'Revenue', 'Profit', 'Units']
                
                # Count transactions per year
                transaction_counts = df.groupby('Year').size().reset_index(name='Transactions')
                yoy_summary = yoy_summary.merge(transaction_counts, on='Year')
                
                yoy_summary['Profit Margin %'] = (yoy_summary['Profit'] / yoy_summary['Revenue'] * 100).round(2)
                yoy_summary = yoy_summary.sort_values('Year')
                
                # Calculate YoY growth
                yoy_summary['Revenue Growth %'] = yoy_summary['Revenue'].pct_change() * 100
                yoy_summary['Profit Growth %'] = yoy_summary['Profit'].pct_change() * 100
                yoy_summary['Units Growth %'] = yoy_summary['Units'].pct_change() * 100
            else:
                yoy_summary = pd.DataFrame()
            
            # YoY by business
            yoy_business = df.groupby(['Year', 'Business'])['Revenue'].sum().reset_index() if 'Year' in df.columns and 'Business' in df.columns else pd.DataFrame()
            if not yoy_business.empty:
                yoy_business.columns = ['Year', 'Business', 'Revenue']
            
            # YoY by brand
            yoy_brand = df.groupby(['Year', 'Brand'])['Revenue'].sum().reset_index() if 'Year' in df.columns and 'Brand' in df.columns else pd.DataFrame()
            if not yoy_brand.empty:
                yoy_brand.columns = ['Year', 'Brand', 'Revenue']
            
            data_dict = {
                'YoY Summary': yoy_summary,
                'YoY by Business': yoy_business,
                'YoY by Brand': yoy_brand
            }
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_data = self._create_excel_with_formatting(data_dict, "Year-over-Year Comparison")
            return file_data, f"yoy_comparison_{timestamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            
        except Exception as e:
            logger.error(f"Error generating YoY comparison report: {str(e)}")
            raise

    async def generate_monthly_trends_report(
        self,
        years: Optional[str] = None,
        businesses: Optional[str] = None
    ) -> Tuple[bytes, str, str]:
        """Generate Monthly Trends Report"""
        try:
            df = await self._get_filtered_data(years=years, businesses=businesses)
            
            if df.empty:
                raise ValueError("No data found for monthly trends")
            
            # Monthly summary
            if 'Year' in df.columns and 'Month_Name' in df.columns:
                monthly_summary = df.groupby(['Year', 'Month_Name']).agg({
                    'Revenue': 'sum',
                    'Gross_Profit': 'sum',
                    'Units': 'sum'
                }).reset_index()
                monthly_summary.columns = ['Year', 'Month', 'Revenue', 'Profit', 'Units']
                
                # Add month order for proper sorting
                month_order = ['January', 'February', 'March', 'April', 'May', 'June', 
                              'July', 'August', 'September', 'October', 'November', 'December']
                monthly_summary['Month_Order'] = monthly_summary['Month'].map({m: i for i, m in enumerate(month_order)})
                monthly_summary = monthly_summary.sort_values(['Year', 'Month_Order'])
                monthly_summary = monthly_summary.drop('Month_Order', axis=1)
            else:
                monthly_summary = pd.DataFrame()
            
            # Monthly by business
            if 'Year' in df.columns and 'Month_Name' in df.columns and 'Business' in df.columns:
                monthly_business = df.groupby(['Year', 'Month_Name', 'Business'])['Revenue'].sum().reset_index()
                monthly_business.columns = ['Year', 'Month', 'Business', 'Revenue']
            else:
                monthly_business = pd.DataFrame()
            
            # Seasonal patterns (average by month across all years)
            if 'Month_Name' in df.columns:
                seasonal = df.groupby('Month_Name').agg({
                    'Revenue': 'mean',
                    'Units': 'mean'
                }).reset_index()
                seasonal.columns = ['Month', 'Avg Revenue', 'Avg Units']
                
                month_order = ['January', 'February', 'March', 'April', 'May', 'June', 
                              'July', 'August', 'September', 'October', 'November', 'December']
                seasonal['Month_Order'] = seasonal['Month'].map({m: i for i, m in enumerate(month_order)})
                seasonal = seasonal.sort_values('Month_Order')
                seasonal = seasonal.drop('Month_Order', axis=1)
            else:
                seasonal = pd.DataFrame()
            
            data_dict = {
                'Monthly Trends': monthly_summary,
                'Monthly by Business': monthly_business,
                'Seasonal Patterns': seasonal
            }
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_data = self._create_excel_with_formatting(data_dict, "Monthly Trends")
            return file_data, f"monthly_trends_{timestamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            
        except Exception as e:
            logger.error(f"Error generating monthly trends report: {str(e)}")
            raise
